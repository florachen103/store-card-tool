"""
Native frontend entry for the work-card tool.
Run with: python app.py
"""

from __future__ import annotations

import base64
import gc
import io
import math
import os
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

import pandas as pd
from flask import Flask, Response, jsonify, redirect, render_template_string, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.exceptions import RequestEntityTooLarge

import browser_pdf_export as browser_export
from excel_image_extractor import extract_images_from_excel


app = Flask(__name__)

MAX_UPLOAD_BYTES = int(os.environ.get("STORE_CARD_MAX_UPLOAD_MB", "20")) * 1024 * 1024
MAX_CARDS_PER_UPLOAD = int(os.environ.get("STORE_CARD_MAX_CARDS", "50"))
ARTIFACT_TTL_SECONDS = int(os.environ.get("STORE_CARD_ARTIFACT_TTL_SECONDS", "1800"))
MAX_ARTIFACTS = int(os.environ.get("STORE_CARD_MAX_ARTIFACTS", "3"))

app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

ARTIFACTS: dict[str, dict] = {}
TMP_PREVIEW_PATH = Path(tempfile.gettempdir()) / "store-card-latest-preview.html"
LOGO_B64 = base64.b64encode((Path(__file__).resolve().parent / "assets" / "logo_live.svg").read_bytes()).decode()
STAR_SRC = browser_export._load_star_b64()
CENTER_BG_B64 = base64.b64encode((Path(__file__).resolve().parent / "assets" / "card_center_bg.png").read_bytes()).decode()
SLOGAN_B64 = base64.b64encode((Path(__file__).resolve().parent / "assets" / "huyou_da_jiankang_calligraphy_vector.svg").read_bytes()).decode()


@app.errorhandler(RequestEntityTooLarge)
def handle_upload_too_large(_exc):
    max_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
    return jsonify({"ok": False, "error": f"Excel 文件不能超过 {max_mb}MB"}), 413


INDEX_HTML = """
<!DOCTYPE html>
<html lang="zh-CN">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>高济药房 · 员工工牌制作系统</title>
  <style>
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: "Noto Sans SC", "PingFang SC", "Microsoft YaHei", sans-serif;
      background: #eff1f5;
      color: #1a1a1a;
    }
    .site-header {
      background: #fff;
      border-bottom: 1px solid #ececec;
    }
    .site-header-inner {
      max-width: 1160px;
      margin: 0 auto;
      padding: 16px 20px;
      display: flex;
      align-items: center;
      gap: 14px;
    }
    .brand-badge {
      width: 44px;
      height: 44px;
      border-radius: 8px;
      background: #e60036;
      color: #fff;
      display: flex;
      align-items: center;
      justify-content: center;
      font-size: 26px;
      font-weight: 700;
    }
    .brand-title {
      font-size: 18px;
      font-weight: 700;
      color: #1a1a1a;
      margin: 0;
    }
    .brand-subtitle {
      margin: 2px 0 0;
      color: #9ca3af;
      font-size: 12px;
      font-weight: 500;
    }
    .page {
      max-width: 1160px;
      margin: 0 auto;
      padding: 20px;
    }
    .step-grid {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 16px;
      margin-bottom: 20px;
    }
    .card {
      background: #fff;
      border-radius: 12px;
      box-shadow: 0 1px 4px rgba(0,0,0,0.08);
    }
    .step-card {
      padding: 16px;
      display: flex;
      gap: 12px;
      align-items: flex-start;
    }
    .step-index {
      width: 26px;
      height: 26px;
      background: #fff0f3;
      border-radius: 50%;
      display: flex;
      align-items: center;
      justify-content: center;
      font-weight: 700;
      color: #E60036;
      font-size: 13px;
      flex-shrink: 0;
    }
    .step-title {
      margin: 0;
      font-weight: 600;
      font-size: 13px;
      color: #1a1a1a;
    }
    .step-desc {
      margin: 4px 0 0;
      font-size: 11px;
      color: #999;
      line-height: 1.55;
    }
    .drop-card {
      margin-bottom: 20px;
      overflow: hidden;
      padding: 0;
    }
    .drop-zone {
      margin: 0;
      border: 2px dashed #E60036;
      border-radius: 16px;
      cursor: pointer;
      transition: background 0.2s, border-color 0.2s;
      background: #fff;
      padding: 48px 24px;
      text-align: center;
    }
    .drop-zone:hover,
    .drop-zone.drag-active {
      background: #fff5f5;
      border-color: #a00020;
    }
    .drop-icon-wrap {
      width: 60px;
      height: 60px;
      background: #fff0f3;
      border-radius: 50%;
      display: flex;
      align-items: center;
      justify-content: center;
      margin: 0 auto 14px;
    }
    .drop-title {
      font-size: 17px;
      font-weight: 600;
      color: #1a1a1a;
      margin: 0;
    }
    .drop-subtitle {
      font-size: 13px;
      color: #999;
      margin: 6px 0 0;
      line-height: 1.6;
    }
    .btn-outline,
    .btn-red {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      border-radius: 999px;
      font-size: 13px;
      font-weight: 500;
      padding: 8px 18px;
      cursor: pointer;
      text-decoration: none;
      border: 1.5px solid #E60036;
    }
    .btn-outline {
      background: #fff;
      color: #E60036;
    }
    .btn-red {
      background: #E60036;
      color: #fff;
    }
    .progress-wrap,
    .error-wrap {
      display: none;
      padding: 16px 28px 20px;
    }
    .progress-row {
      display: flex;
      justify-content: space-between;
      margin-bottom: 8px;
      font-size: 13px;
    }
    .progress-track {
      background: #f0f0f0;
      border-radius: 4px;
      height: 6px;
    }
    .progress-bar {
      background: #E60036;
      height: 6px;
      border-radius: 4px;
      width: 0;
      transition: width 0.2s;
    }
    .progress-detail {
      font-size: 11px;
      color: #aaa;
      margin-top: 6px;
    }
    .error-box {
      background: #fff0f3;
      border: 1px solid #fcc;
      border-radius: 10px;
      padding: 14px 16px;
      display: flex;
      gap: 10px;
    }
    .preview-card {
      display: none;
    }
    .preview-toolbar {
      padding: 14px 18px;
      margin-bottom: 14px;
      display: flex;
      align-items: center;
      justify-content: space-between;
      flex-wrap: wrap;
      gap: 12px;
    }
    .preview-title {
      font-weight: 700;
      font-size: 15px;
      color: #1a1a1a;
    }
    .preview-meta {
      font-size: 13px;
      color: #999;
    }
    .template-controls {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      align-items: center;
    }
    .template-control {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      color: #555;
      font-size: 12px;
      font-weight: 600;
    }
    .choice-control {
      position: relative;
    }
    .choice-trigger {
      height: 32px;
      border: 1px solid #ddd;
      border-radius: 8px;
      background: #fff;
      color: #1a1a1a;
      font-size: 13px;
      font-weight: 600;
      padding: 0 30px 0 10px;
      cursor: pointer;
      min-width: 76px;
      text-align: left;
      position: relative;
    }
    .choice-trigger::after {
      content: '';
      position: absolute;
      right: 10px;
      top: 50%;
      width: 7px;
      height: 7px;
      border-right: 1.5px solid #555;
      border-bottom: 1.5px solid #555;
      transform: translateY(-65%) rotate(45deg);
    }
    .choice-menu {
      display: none;
      position: absolute;
      top: calc(100% + 6px);
      left: 0;
      z-index: 1000;
      min-width: 100%;
      padding: 6px;
      background: #fff;
      border: 1px solid #e5e7eb;
      border-radius: 12px;
      box-shadow: 0 12px 28px rgba(15, 23, 42, 0.16);
    }
    .choice-control.open .choice-menu {
      display: grid;
      gap: 3px;
    }
    .choice-option {
      border: 0;
      background: transparent;
      color: #1a1a1a;
      border-radius: 8px;
      height: 30px;
      padding: 0 28px 0 10px;
      text-align: left;
      font-size: 13px;
      font-weight: 600;
      white-space: nowrap;
      cursor: pointer;
      position: relative;
    }
    .choice-option:hover {
      background: #fff0f3;
      color: #E60036;
    }
    .choice-option.is-selected {
      background: #E60036;
      color: #fff;
    }
    .choice-option.is-selected::before {
      content: '';
      position: absolute;
      left: 10px;
      top: 50%;
      width: 6px;
      height: 11px;
      border-right: 2px solid #fff;
      border-bottom: 2px solid #fff;
      transform: translateY(-62%) rotate(45deg);
    }
    .choice-option.is-selected {
      padding-left: 28px;
    }
    .print-tip {
      background: #fffbeb;
      border: 1px solid #fde68a;
      border-radius: 10px;
      padding: 10px 14px;
      margin-bottom: 14px;
      font-size: 12px;
      color: #92400e;
      line-height: 1.6;
    }
    .preview-shell {
      overflow: hidden;
      padding: 18px;
      background: #eef0f4;
    }
    .badge-preview-grid {
      display: flex;
      flex-direction: column;
      gap: 18px;
      align-items: center;
      justify-content: center;
      overflow-x: auto;
      padding-bottom: 4px;
    }
    .preview-a4-wrap {
      width: max-content;
    }
    .preview-a4-label {
      margin: 0 0 8px;
      color: #6b7280;
      font-size: 12px;
      font-weight: 600;
    }
    .preview-a4-page {
      width: 210mm;
      height: 297mm;
      padding: 5mm 7mm;
      background: #fff;
      box-shadow: 0 10px 28px rgba(15, 23, 42, 0.12);
      box-sizing: border-box;
    }
    .preview-a4-grid {
      display: grid;
      grid-template-columns: 86mm 86mm;
      gap: 2mm 4mm;
      width: 176mm;
      margin: 0 auto;
    }
    .preview-a4-page .badge-card {
      box-shadow: none;
      border: 0.3mm solid #e0e0e0;
    }
    .actions {
      margin-top: 14px;
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
    }
    .footer {
      text-align: center;
      padding: 16px 0 24px;
      color: #bbb;
      font-size: 12px;
    }
    .print-area { display: none; }
    .badge-card {
      width: 86mm;
      height: 54mm;
      background: #fff;
      overflow: hidden;
      display: flex;
      flex-direction: column;
      position: relative;
      flex-shrink: 0;
      box-shadow: 0 2px 8px rgba(0,0,0,0.18);
    }
    .badge-card.staff-card .bh-role {
      display: none;
    }
    .badge-card.staff-card .bh-logo {
      left: 3mm;
    }
    .badge-card-bg {
      position: absolute;
      left: 50%;
      top: 50%;
      width: 45.5mm;
      height: auto;
      transform: translate(calc(-50% - 5px), calc(-50% + 10px));
      opacity: 0.56;
      pointer-events: none;
      z-index: 0;
    }
    .bh {
      background: #E50036;
      height: 11.3mm;
      margin-top: -5px;
      display: flex;
      align-items: center;
      padding: 0 2mm 0 0;
      gap: 2mm;
      flex-shrink: 0;
      position: relative;
      overflow: visible;
    }
    .bh-role {
      display: flex;
      flex-direction: column;
      align-items: center;
      flex-shrink: 0;
      align-self: flex-start;
      width: 10.8mm;
      border-radius: 2mm 2mm 0 0;
      margin-left: 12px;
      margin-top: 10px;
      overflow: visible;
    }
    .bh-role-top {
      width: 10.8mm;
      height: 8.6mm;
      background: #fff;
      border-radius: 2mm 2mm 0 0;
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      gap: 0.8mm;
      flex-shrink: 0;
    }
    .bh-role-bottom {
      width: 10.8mm;
      background: #E50036;
      border-radius: 0 0 2mm 2mm;
      display: flex;
      align-items: center;
      justify-content: center;
      padding: 0.8mm 0.5mm;
    }
    .bh-role-crown svg { width: 4mm; height: 2.8mm; display: block; }
    .bh-role-cn {
      font-size: 3.5mm;
      font-weight: 900;
      color: #E50036;
      line-height: 1;
      text-align: center;
      letter-spacing: 0.5mm;
    }
    .bh-role-en {
      font-size: 1.1mm;
      font-weight: 700;
      color: #fff;
      line-height: 1;
      text-align: center;
      white-space: nowrap;
      letter-spacing: 0.1mm;
    }
    .bh-logo {
      position: absolute;
      left: 14.2mm;
      top: 10px;
      height: 6.55mm;
      width: auto;
      object-fit: contain;
      z-index: 200;
    }
    .bh-deco {
      position: absolute;
      right: 0;
      top: 0;
      height: 11.3mm;
      width: 24mm;
      display: flex;
      align-items: center;
      justify-content: flex-end;
      gap: 2mm;
      overflow: visible;
      pointer-events: none;
    }
    .bh-deco-circle {
      width: calc(14.5mm + 58px);
      height: calc(17mm - 5px);
      background: #fff;
      border-radius: 999px;
      flex-shrink: 0;
      transform: translate(calc(4mm + 18px), 3px);
      position: relative;
      z-index: 10;
    }
    .bh-deco-circle::before {
      content: '';
      position: absolute;
      inset: calc(3.6mm - 3px);
      background: #E50036;
      border-radius: inherit;
    }
    .bh-deco-circle::after {
      content: '';
      position: absolute;
      top: 0;
      bottom: 0;
      left: 50%;
      transform: translateX(-50%);
      width: 11px;
      background: #fff;
    }
    .bh-deco-pill {
      width: 5.5mm;
      height: 17mm;
      background: #fff;
      border-radius: 3mm;
      flex-shrink: 0;
      transform: translateX(4mm);
      position: relative;
    }
    .bh-deco-pill::before {
      content: '';
      position: absolute;
      inset: 3.6mm;
      background: #E50036;
      border-radius: calc(3mm - 1mm);
    }
    .bb {
      flex: 1;
      display: grid;
      grid-template-columns: 25mm 1fr 22mm;
      grid-template-rows: 11.2mm 1fr;
      padding: 2.8mm 2.1mm 2.2mm 2.1mm;
      column-gap: 1.6mm;
      row-gap: 1.2mm;
      background: transparent;
      min-height: 0;
      position: relative;
      z-index: 2;
    }
    .bb-photo {
      grid-column: 1;
      grid-row: 1 / 3;
      padding-top: calc(2.15mm + 3px);
      padding-left: 4px;
      display: flex;
      flex-direction: column;
    }
    .bb-photo-frame {
      width: 25mm;
      height: 35mm;
      flex-shrink: 0;
      border: 0.5mm dashed #999;
      position: relative;
      display: flex;
      align-items: center;
      justify-content: center;
      overflow: hidden;
    }
    .bb-photo-frame img {
      position: absolute;
      inset: 0;
      width: 100%;
      height: 100%;
      object-fit: cover;
      display: block;
    }
    .bb-qr-frame img {
      position: absolute;
      inset: 0;
      width: 100%;
      height: 100%;
      object-fit: contain;
      display: block;
    }
    .bb-photo-frame svg.placeholder-x,
    .bb-qr-frame svg.placeholder-x {
      position: absolute;
      inset: 0;
      width: 100%;
      height: 100%;
    }
    .bb-photo-frame .ph-text {
      position: absolute;
      font-size: 2.8mm;
      color: #aaa;
      font-weight: 500;
      writing-mode: vertical-rl;
      text-orientation: mixed;
      letter-spacing: 1mm;
    }
    .bb-stars {
      grid-column: 2 / 4;
      grid-row: 1;
      display: flex;
      justify-content: center;
      align-items: center;
      gap: 2px;
    }
    .bb-stars.slogan-mode {
      justify-content: center;
      padding: 0 2mm;
    }
    .bb-star {
      width: 9.8mm;
      height: 9.8mm;
      object-fit: contain;
      flex-shrink: 0;
    }
    .bb-slogan-img {
      width: 45mm;
      max-height: 8.6mm;
      object-fit: contain;
      display: block;
    }
    .bb-fields {
      grid-column: 2;
      grid-row: 2;
      display: flex;
      flex-direction: column;
      gap: 0;
      align-self: end;
      padding-right: 4%;
      margin-left: 4px;
    }
    .bb-field {
      display: flex;
      align-items: flex-end;
      margin-bottom: 1.8mm;
      gap: 0;
    }
    .bb-field-label {
      font-size: 3.65mm;
      font-weight: 700;
      color: #1a1a1a;
      white-space: nowrap;
      flex-shrink: 0;
      line-height: 1.2;
    }
    .bb-field-line {
      flex: 1;
      border-bottom: 0.5mm solid #1a1a1a;
      min-width: 6mm;
      margin-left: 0.7mm;
      padding-bottom: 0.35mm;
      font-size: 3.45mm;
      font-weight: 600;
      color: #1a1a1a;
      line-height: 1.2;
      white-space: nowrap;
      overflow: hidden;
      text-overflow: ellipsis;
    }
    .bb-qr {
      grid-column: 3;
      grid-row: 2;
      align-self: end;
    }
    .bb-qr-frame {
      width: 22mm;
      height: 22mm;
      transform: translate(-4px, 3px);
      border: 0.5mm dashed #999;
      position: relative;
      display: flex;
      align-items: center;
      justify-content: center;
      overflow: hidden;
    }
    .bb-qr-frame .ph-text {
      position: absolute;
      font-size: 2.8mm;
      color: #aaa;
      font-weight: 500;
    }
    .bb-qr-label {
      position: absolute;
      bottom: 1.2mm;
      left: 0;
      right: 0;
      text-align: center;
      font-size: 2.2mm;
      color: #888;
      pointer-events: none;
    }
    .bt {
      height: 2.9mm;
      background: #E60036;
      flex-shrink: 0;
    }
    @media print {
      @page { size: A4 portrait; margin: 5mm 7mm 5mm 7mm; }
      html, body { margin: 0; padding: 0; background: white !important; }
      .no-print { display: none !important; }
      .print-area { display: block !important; }
      .site-header,
      .footer,
      .drop-card,
      .preview-card {
        display: none !important;
      }
      .page { display: none !important; }
      .badge-page {
        display: grid !important;
        grid-template-columns: 86mm 86mm;
        gap: 2mm 4mm;
        width: 176mm;
        margin: 0 auto;
        padding: 0;
      }
      .badge-card {
        box-shadow: none !important;
        border: 0.3mm solid #e0e0e0;
        break-inside: avoid;
      }
      * {
        -webkit-print-color-adjust: exact !important;
        print-color-adjust: exact !important;
        color-adjust: exact !important;
      }
      .bb-photo { padding-top: calc(2.15mm - 3px) !important; }
    }
    @media (max-width: 900px) {
      .step-grid { grid-template-columns: 1fr; }
      .badge-preview-grid {
        align-items: flex-start;
        justify-content: flex-start;
      }
    }
  </style>
</head>
<body>
  <div class="print-area" id="printArea"></div>
  <header class="site-header">
    <div class="site-header-inner">
      <div class="brand-badge">高</div>
      <div>
        <p class="brand-title">员工工牌制作系统</p>
        <p class="brand-subtitle">高济药房 · CBWELL PHARMACY</p>
      </div>
    </div>
  </header>

  <main class="page">
    <section class="step-grid">
      <div class="card step-card">
        <div class="step-index">1</div>
        <div>
          <p class="step-title">准备 Excel 表格</p>
          <p class="step-desc">列1姓名·列2职务·列3工号·列4嵌入照片·列5嵌入企微码</p>
        </div>
      </div>
      <div class="card step-card">
        <div class="step-index">2</div>
        <div>
          <p class="step-title">导入 .xlsx 文件</p>
          <p class="step-desc">拖拽或点击上传，自动解析文字 + 嵌入图片</p>
        </div>
      </div>
      <div class="card step-card">
        <div class="step-index">3</div>
        <div>
          <p class="step-title">打印输出</p>
          <p class="step-desc">A4纸·每页10张(2×5)·实际大小·关闭页眉页脚</p>
        </div>
      </div>
    </section>

    <section class="card drop-card" id="uploadCard">
      <div class="drop-zone" id="dropZone">
        <input type="file" id="fileInput" accept=".xlsx" style="display:none">
        <div class="drop-icon-wrap">
          <svg width="28" height="28" fill="none" stroke="#E60036" stroke-width="1.5" viewBox="0 0 24 24">
            <path stroke-linecap="round" stroke-linejoin="round" d="M9 12h6m-6 4h6m2 5H7a2 2 0 01-2-2V5a2 2 0 012-2h5.586a1 1 0 01.707.293l5.414 5.414a1 1 0 01.293.707V19a2 2 0 01-2 2z"/>
          </svg>
        </div>
        <p class="drop-title">点击或拖拽上传 Excel 文件</p>
        <p class="drop-subtitle">支持 .xlsx 格式，自动提取照片和企业微信二维码</p>
        <div style="margin-top:12px;display:inline-flex;gap:10px;flex-wrap:wrap;justify-content:center;">
          <a class="btn-outline" href="/template.xlsx" id="downloadTemplateBtn">
            <svg width="15" height="15" fill="none" stroke="currentColor" stroke-width="2" viewBox="0 0 24 24">
              <path stroke-linecap="round" stroke-linejoin="round" d="M4 16v2a2 2 0 002 2h12a2 2 0 002-2v-2M7 10l5 5 5-5M12 15V3"/>
            </svg>
            下载 Excel 模板
          </a>
          <button class="btn-red" id="uploadTemplateBtn" type="button">
            <svg width="15" height="15" fill="none" stroke="currentColor" stroke-width="2" viewBox="0 0 24 24">
              <path stroke-linecap="round" stroke-linejoin="round" d="M12 21V9m0 0l-5 5m5-5l5 5M5 3h14a2 2 0 012 2v2H3V5a2 2 0 012-2z"/>
            </svg>
            上传 Excel 模板
          </button>
        </div>
      </div>

      <div class="progress-wrap" id="progressWrap">
        <div class="progress-row">
          <span id="progText" style="font-weight:500;color:#333;">解析中...</span>
          <span id="progPct" style="color:#999;">0%</span>
        </div>
        <div class="progress-track"><div class="progress-bar" id="progBar"></div></div>
        <p class="progress-detail" id="progDetail"></p>
      </div>

      <div class="error-wrap" id="errorWrap">
        <div class="error-box">
          <svg width="18" height="18" fill="#ef4444" viewBox="0 0 20 20" style="flex-shrink:0;margin-top:1px;">
            <path fill-rule="evenodd" d="M10 18a8 8 0 100-16 8 8 0 000 16zM8.707 7.293a1 1 0 00-1.414 1.414L8.586 10l-1.293 1.293a1 1 0 101.414 1.414L10 11.414l1.293 1.293a1 1 0 001.414-1.414L11.414 10l1.293-1.293a1 1 0 00-1.414-1.414L10 8.586 8.707 7.293z" clip-rule="evenodd"/>
          </svg>
          <div>
            <p style="font-weight:600;font-size:13px;color:#c00;margin:0;">解析失败</p>
            <p id="errorMsg" style="font-size:12px;color:#e55;margin:2px 0 0;"></p>
          </div>
        </div>
      </div>

      <div id="previewCard" class="preview-card">
      <div class="card preview-toolbar">
        <div style="display:flex;align-items:center;gap:16px;flex-wrap:wrap;">
          <span class="preview-title">工牌预览</span>
          <span class="preview-meta">共 <b style="color:#E60036;" id="totalNum">0</b> 张 · <b style="color:#555;" id="pageNum">0</b> 页</span>
        </div>
        <div class="template-controls">
          <label class="template-control">
            角色
            <div class="choice-control" data-choice="role">
              <button class="choice-trigger" type="button" data-choice-trigger>店长</button>
              <div class="choice-menu">
                <button class="choice-option is-selected" type="button" data-value="manager">店长</button>
                <button class="choice-option" type="button" data-value="staff">店员</button>
              </div>
            </div>
          </label>
          <label class="template-control">
            星级
            <div class="choice-control" data-choice="stars">
              <button class="choice-trigger" type="button" data-choice-trigger>五星</button>
              <div class="choice-menu">
                <button class="choice-option is-selected" type="button" data-value="5">五星</button>
                <button class="choice-option" type="button" data-value="4">四星</button>
                <button class="choice-option" type="button" data-value="3">三星</button>
                <button class="choice-option" type="button" data-value="2">二星</button>
                <button class="choice-option" type="button" data-value="1">一星</button>
                <button class="choice-option" type="button" data-value="0">无星</button>
              </div>
            </div>
          </label>
          <button class="btn-outline" id="reuploadBtn" type="button">重新上传</button>
          <button class="btn-red" id="printBtn" type="button">打印工牌</button>
        </div>
      </div>

      <div class="print-tip">
        <b>打印设置：</b>A4纸 · 实际大小(100%) · 关闭页眉页脚 · 不勾选“适应页面” · 每页自动排列 10 张工牌（2列×5行）· 挂卡尺寸 86×54mm
      </div>

      <div class="card preview-shell">
        <div class="badge-preview-grid" id="badgeGrid"></div>
      </div>

      </div>
    </section>
  </main>

  <div class="footer">© 2026 高济药房 · 员工工牌制作系统 · 保留所有权利</div>

  <script>
    const fileInput = document.getElementById('fileInput');
    const dropZone = document.getElementById('dropZone');
    const progressWrap = document.getElementById('progressWrap');
    const errorWrap = document.getElementById('errorWrap');
    const errorMsg = document.getElementById('errorMsg');
    const previewCard = document.getElementById('previewCard');
    const badgeGrid = document.getElementById('badgeGrid');
    const printArea = document.getElementById('printArea');
    const progText = document.getElementById('progText');
    const progPct = document.getElementById('progPct');
    const progBar = document.getElementById('progBar');
    const progDetail = document.getElementById('progDetail');
    const totalNum = document.getElementById('totalNum');
    const pageNum = document.getElementById('pageNum');
    const printBtn = document.getElementById('printBtn');
    const reuploadBtn = document.getElementById('reuploadBtn');
    const uploadTemplateBtn = document.getElementById('uploadTemplateBtn');

    let currentArtifact = null;
    let employees = [];
    let selectedRole = 'manager';
    let selectedStarCount = 5;
    const LOGO_B64 = "__LOGO_B64__";
    const STAR_SRC = "__STAR_SRC__";
    const CENTER_BG_B64 = "__CENTER_BG_B64__";
    const SLOGAN_SRC = "data:image/svg+xml;base64,__SLOGAN_B64__";
    const POS_EN = {
      '店长':'Store Manager','副店长':'Asst. Manager','店员':'Store Staff',
      '药师':'Pharmacist','执业药师':'Pharmacist','收银员':'Cashier',
      '主任':'Director','经理':'Manager','顾问':'Consultant'
    };

    function getPosEn(pos) { return POS_EN[pos] || 'Employee'; }

    function esc(s) {
      return String(s || '').replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');
    }

    function setProgress(pct, text, detail='') {
      progressWrap.style.display = 'block';
      progPct.textContent = pct + '%';
      progBar.style.width = pct + '%';
      progText.textContent = text;
      progDetail.textContent = detail;
    }

    function hideError() {
      errorWrap.style.display = 'none';
      errorMsg.textContent = '';
    }

    function showError(msg) {
      errorWrap.style.display = 'block';
      errorMsg.textContent = msg;
      progressWrap.style.display = 'none';
    }

    function resetPreview() {
      dropZone.style.display = 'block';
      previewCard.style.display = 'none';
      badgeGrid.innerHTML = '';
      printArea.innerHTML = '';
      employees = [];
      currentArtifact = null;
    }

    function buildStarsHtml() {
      if (selectedStarCount <= 0) {
        return `<img class="bb-slogan-img" src="${SLOGAN_SRC}" alt="护佑大众健康">`;
      }
      return Array.from({ length: selectedStarCount }, () => `<img class="bb-star" src="${STAR_SRC}" alt="★">`).join('');
    }

    function makeBadge(emp) {
      const isManager = selectedRole === 'manager';
      const roleCnDisplay = '店长';
      const roleEnDisplay = 'Store Manager';
      const starsHtml = buildStarsHtml();
      const starsClass = selectedStarCount <= 0 ? 'bb-stars slogan-mode' : 'bb-stars';
      const photoHtml = emp.photo
        ? `<img src="${emp.photo}" alt="照片">`
        : `<svg class="placeholder-x" viewBox="0 0 100 100" preserveAspectRatio="none">
            <line x1="0" y1="0" x2="100" y2="100" stroke="#ccc" stroke-width="1.2"/>
            <line x1="100" y1="0" x2="0" y2="100" stroke="#ccc" stroke-width="1.2"/>
          </svg><span class="ph-text">一寸照</span>`;
      const qrHtml = emp.qrcode
        ? `<img src="${emp.qrcode}" alt="二维码">`
        : `<svg class="placeholder-x" viewBox="0 0 100 100" preserveAspectRatio="none">
            <line x1="0" y1="0" x2="100" y2="100" stroke="#ccc" stroke-width="1.2"/>
            <line x1="100" y1="0" x2="0" y2="100" stroke="#ccc" stroke-width="1.2"/>
          </svg><span class="ph-text">企微码</span>`;

      const div = document.createElement('div');
      div.className = isManager ? 'badge-card' : 'badge-card staff-card';
      div.innerHTML = `
        <img class="badge-card-bg" src="data:image/png;base64,${CENTER_BG_B64}" alt="">
        <div class="bh">
          <div class="bh-role">
            <div class="bh-role-top">
              <div class="bh-role-crown">
                <svg viewBox="0 0 56 40" xmlns="http://www.w3.org/2000/svg" fill="none" stroke="#E60036" stroke-linejoin="round" stroke-linecap="round">
                  <rect x="6" y="28" width="44" height="8" rx="2" stroke-width="2.5"/>
                  <path d="M6 28 L6 16 L18 24 L28 6 L38 24 L50 16 L50 28 Z" stroke-width="2.5"/>
                  <circle cx="28" cy="6" r="3.5" stroke-width="2"/>
                  <circle cx="6" cy="16" r="2.5" stroke-width="2"/>
                  <circle cx="50" cy="16" r="2.5" stroke-width="2"/>
                </svg>
              </div>
              <div class="bh-role-cn">${esc(roleCnDisplay)}</div>
            </div>
            <div class="bh-role-bottom"><div class="bh-role-en">${esc(roleEnDisplay)}</div></div>
          </div>
          <img class="bh-logo" src="data:image/svg+xml;base64,${LOGO_B64}" alt="logo">
          <div class="bh-deco">
            <div class="bh-deco-circle"></div>
            <div class="bh-deco-pill"></div>
          </div>
        </div>
        <div class="bb">
          <div class="bb-photo">
            <div class="bb-photo-frame" ${emp.photo ? 'style="border-color:transparent"' : ''}>${photoHtml}</div>
          </div>
          <div class="${starsClass}">${starsHtml}</div>
          <div class="bb-fields">
            <div class="bb-field">
              <span class="bb-field-label">姓名:</span>
              <span class="bb-field-line">${esc(emp.name)}</span>
            </div>
            <div class="bb-field">
              <span class="bb-field-label">职务:</span>
              <span class="bb-field-line">${esc(emp.position)}</span>
            </div>
            <div class="bb-field" style="margin-bottom:0;">
              <span class="bb-field-label">工号:</span>
              <span class="bb-field-line">${esc(emp.employeeId)}</span>
            </div>
          </div>
          <div class="bb-qr">
            <div class="bb-qr-frame" ${emp.qrcode ? 'style="border-color:transparent"' : ''}>
              ${qrHtml}
              ${!emp.qrcode ? '<div class="bb-qr-label">企微码</div>' : ''}
            </div>
          </div>
        </div>
        <div class="bt"></div>`;
      return div;
    }

    function renderAll(emps) {
      badgeGrid.innerHTML = '';
      const perPage = 10;
      const pages = Math.ceil(emps.length / perPage);
      for (let p = 0; p < pages; p++) {
        const pageWrap = document.createElement('div');
        pageWrap.className = 'preview-a4-wrap';

        const pageLabel = document.createElement('p');
        pageLabel.className = 'preview-a4-label';
        pageLabel.textContent = `第 ${p + 1} 页 / 共 ${pages} 页`;

        const pageSheet = document.createElement('div');
        pageSheet.className = 'preview-a4-page';

        const pageGrid = document.createElement('div');
        pageGrid.className = 'preview-a4-grid';

        const slice = emps.slice(p * perPage, (p + 1) * perPage);
        slice.forEach(e => pageGrid.appendChild(makeBadge(e)));
        for (let i = slice.length; i < perPage; i++) {
          const blank = document.createElement('div');
          blank.className = 'badge-card';
          blank.style.visibility = 'hidden';
          pageGrid.appendChild(blank);
        }

        pageSheet.appendChild(pageGrid);
        pageWrap.appendChild(pageLabel);
        pageWrap.appendChild(pageSheet);
        badgeGrid.appendChild(pageWrap);
      }
      totalNum.textContent = emps.length;
      pageNum.textContent = pages;
    }

    function doPrint() {
      printArea.innerHTML = '';
      const perPage = 10;
      const pages = Math.ceil(employees.length / perPage);
      for (let p = 0; p < pages; p++) {
        const pageDiv = document.createElement('div');
        pageDiv.className = 'badge-page';
        if (p < pages - 1) pageDiv.style.pageBreakAfter = 'always';
        const slice = employees.slice(p * perPage, (p + 1) * perPage);
        slice.forEach(e => pageDiv.appendChild(makeBadge(e)));
        for (let i = slice.length; i < perPage; i++) {
          const blank = document.createElement('div');
          blank.className = 'badge-card';
          blank.style.visibility = 'hidden';
          pageDiv.appendChild(blank);
        }
        printArea.appendChild(pageDiv);
      }
      window.print();
    }

    async function processFile(file) {
      if (!file || !file.name.toLowerCase().endsWith('.xlsx')) {
        showError('请上传 .xlsx 格式的 Excel 文件');
        return;
      }

      hideError();
      resetPreview();
      setProgress(10, '上传中...', '正在提交 Excel 文件');

      const formData = new FormData();
      formData.append('file', file);

      try {
        setProgress(35, '解析中...', '正在提取文字和内嵌图片');
        const res = await fetch('/api/upload', { method: 'POST', body: formData });
        const data = await res.json();
        if (!res.ok || !data.ok) {
          throw new Error(data.error || '处理失败');
        }

        setProgress(100, '生成完成', '正在加载预览');
        currentArtifact = data.artifactId;
        employees = data.employees || [];
        renderAll(employees);
        dropZone.style.display = 'none';
        printBtn.onclick = () => doPrint();
        previewCard.style.display = 'block';
        setTimeout(() => { progressWrap.style.display = 'none'; }, 400);
      } catch (err) {
        showError(err.message || '处理失败');
      }
    }

    dropZone.addEventListener('click', () => fileInput.click());
    dropZone.addEventListener('dragover', (e) => {
      e.preventDefault();
      dropZone.classList.add('drag-active');
    });
    dropZone.addEventListener('dragleave', () => dropZone.classList.remove('drag-active'));
    dropZone.addEventListener('drop', (e) => {
      e.preventDefault();
      dropZone.classList.remove('drag-active');
      const file = e.dataTransfer.files && e.dataTransfer.files[0];
      processFile(file);
    });
    fileInput.addEventListener('change', (e) => processFile(e.target.files[0]));
    reuploadBtn.addEventListener('click', () => fileInput.click());
    document.querySelectorAll('[data-choice]').forEach((choice) => {
      const trigger = choice.querySelector('[data-choice-trigger]');
      const options = choice.querySelectorAll('.choice-option');
      trigger.addEventListener('click', (event) => {
        event.stopPropagation();
        document.querySelectorAll('.choice-control.open').forEach((openChoice) => {
          if (openChoice !== choice) openChoice.classList.remove('open');
        });
        choice.classList.toggle('open');
      });
      options.forEach((option) => {
        option.addEventListener('click', (event) => {
          event.stopPropagation();
          options.forEach((item) => item.classList.remove('is-selected'));
          option.classList.add('is-selected');
          trigger.textContent = option.textContent;
          choice.classList.remove('open');
          if (choice.dataset.choice === 'role') {
            selectedRole = option.dataset.value;
          } else {
            selectedStarCount = Number(option.dataset.value);
          }
          if (employees.length) renderAll(employees);
        });
      });
    });
    document.addEventListener('click', () => {
      document.querySelectorAll('.choice-control.open').forEach((choice) => choice.classList.remove('open'));
    });
    document.getElementById('downloadTemplateBtn').addEventListener('click', (e) => e.stopPropagation());
    uploadTemplateBtn.addEventListener('click', (e) => {
      e.stopPropagation();
      fileInput.click();
    });
  </script>
</body>
</html>
"""


def prune_artifacts(max_age_seconds: int = ARTIFACT_TTL_SECONDS) -> None:
    cutoff = time.time() - max_age_seconds
    stale = [key for key, value in ARTIFACTS.items() if value["created_at"] < cutoff]
    if len(ARTIFACTS) - len(stale) > MAX_ARTIFACTS:
        active = sorted(
            ((key, value["created_at"]) for key, value in ARTIFACTS.items() if key not in stale),
            key=lambda item: item[1],
        )
        stale.extend(key for key, _ in active[: max(0, len(ARTIFACTS) - len(stale) - MAX_ARTIFACTS)])
    for key in stale:
        artifact = ARTIFACTS.pop(key, None)
        if artifact and artifact.get("tmpdir"):
            shutil.rmtree(artifact["tmpdir"], ignore_errors=True)


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "员工工牌模板"

    headers = ["姓名", "职务", "工号", "照片", "企业微信二维码"]
    samples = [
        ["李美华", "店长", "GJ-2024001", "在该列单元格内插入照片", "在该列单元格内插入企微码"],
        ["张晨", "副店长", "GJ-2024002", "在该列单元格内插入照片", "在该列单元格内插入企微码"],
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E33143")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, row in enumerate(samples, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 24

    guide = wb.create_sheet("填写说明")
    guide["A1"] = "使用说明"
    guide["A1"].font = Font(bold=True, size=14)
    instructions = [
        "1. 第一行表头请保持不变：姓名、职务、工号、照片、企业微信二维码",
        "2. 照片请直接插入到“照片”列对应单元格内",
        "3. 企微二维码请直接插入到“企业微信二维码”列对应单元格内",
        "4. 保存为 .xlsx 后上传到系统",
    ]
    for idx, text in enumerate(instructions, start=3):
        guide[f"A{idx}"] = text
    guide.column_dimensions["A"].width = 72

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pil_to_b64(img, quality: int = 82) -> str:
    if img is None:
        return ""
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="JPEG", quality=quality, optimize=True)
    return base64.b64encode(buf.getvalue()).decode()


def parse_excel(file_bytes: bytes) -> tuple[list[dict], dict]:
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, nrows=MAX_CARDS_PER_UPLOAD + 1).fillna("")
    if len(df) > MAX_CARDS_PER_UPLOAD:
        raise ValueError(f"单次最多生成 {MAX_CARDS_PER_UPLOAD} 张工牌，请拆分 Excel 后再上传")

    max_excel_row = len(df) + 1
    photo_by_row, qr_by_row = extract_images_from_excel(file_bytes, max_row_1based=max_excel_row)

    warnings = {"missing_photo": [], "missing_qr": []}
    staff_list = []
    for i, row in df.iterrows():
        excel_row = i + 2
        name = row.get("姓名", f"第{excel_row}行")
        if excel_row not in photo_by_row:
            warnings["missing_photo"].append(f"行{excel_row}（{name}）")
        if excel_row not in qr_by_row:
            warnings["missing_qr"].append(f"行{excel_row}（{name}）")
        staff_list.append(
            {
                "name": row.get("姓名", ""),
                "title": row.get("职务", ""),
                "staff_id": row.get("工号", ""),
                "photo": photo_by_row.get(excel_row),
                "qr_code": qr_by_row.get(excel_row),
            }
        )
    return staff_list, warnings


def build_browser_cards(staff_list: list[dict]) -> list[dict]:
    cards = []
    for staff in staff_list:
        photo_b64 = pil_to_b64(staff.get("photo"), quality=82)
        qr_b64 = pil_to_b64(staff.get("qr_code"), quality=88)
        cards.append(
            {
                "name": staff["name"],
                "title": staff["title"],
                "staff_id": staff["staff_id"],
                "photo_b64": photo_b64,
                "qr_b64": qr_b64,
            }
        )
    return cards


def build_client_employees(cards: list[dict]) -> list[dict]:
    items = []
    for card in cards:
        photo_b64 = card.get("photo_b64", "")
        qr_b64 = card.get("qr_b64", "")
        items.append(
            {
                "name": card["name"],
                "position": card["title"],
                "employeeId": card["staff_id"],
                "photo": f"data:image/jpeg;base64,{photo_b64}" if photo_b64 else "",
                "qrcode": f"data:image/jpeg;base64,{qr_b64}" if qr_b64 else "",
            }
        )
    return items


def build_png_zip(pdf_path: Path, total_pages: int, output_path: Path) -> None:
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("未找到 pdftoppm，无法生成 PNG ZIP")

    with tempfile.TemporaryDirectory() as tmpdir:
        out_prefix = Path(tmpdir) / "page"
        cmd = [
            pdftoppm,
            "-png",
            "-r",
            "180",
            "-f",
            "1",
            "-l",
            str(total_pages),
            str(pdf_path),
            str(out_prefix),
        ]
        try:
            subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("PNG 导出超时，请减少单次工牌数量后重试") from exc
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for idx in range(1, total_pages + 1):
                page_path = Path(f"{out_prefix}-{idx}.png")
                if not page_path.exists():
                    break
                zf.write(page_path, f"page_{idx:02d}.png")


def get_artifact(artifact_id: str) -> dict:
    artifact = ARTIFACTS.get(artifact_id)
    if not artifact:
        raise KeyError("未找到预览内容，请重新上传 Excel 文件")
    return artifact


@app.get("/")
def index():
    html = (
        INDEX_HTML.replace("__LOGO_B64__", LOGO_B64)
        .replace("__STAR_SRC__", STAR_SRC)
        .replace("__CENTER_BG_B64__", CENTER_BG_B64)
        .replace("__SLOGAN_B64__", SLOGAN_B64)
    )
    return render_template_string(html)


@app.get("/template.xlsx")
def download_template():
    data = build_template_xlsx()
    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name="工牌信息模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/upload")
def upload_excel():
    prune_artifacts()
    artifact_tmpdir = None
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"ok": False, "error": "请选择 Excel 文件"}), 400
    if not file.filename.lower().endswith(".xlsx"):
        return jsonify({"ok": False, "error": "请上传 .xlsx 格式的 Excel 文件"}), 400

    try:
        if request.content_length and request.content_length > MAX_UPLOAD_BYTES:
            max_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
            return jsonify({"ok": False, "error": f"Excel 文件不能超过 {max_mb}MB"}), 413

        file_bytes = file.read()
        if len(file_bytes) > MAX_UPLOAD_BYTES:
            max_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
            return jsonify({"ok": False, "error": f"Excel 文件不能超过 {max_mb}MB"}), 413

        staff_list, warnings = parse_excel(file_bytes)
        cards = build_browser_cards(staff_list)
        client_employees = build_client_employees(cards)
        del staff_list
        total_cards = len(cards)
        total_pages = max(1, math.ceil(total_cards / browser_export.CARDS_PER_PAGE))

        if total_cards > MAX_CARDS_PER_UPLOAD:
            return jsonify({"ok": False, "error": f"单次最多生成 {MAX_CARDS_PER_UPLOAD} 张工牌"}), 400

        html_content = browser_export.build_pdf_html(cards)
        preview_html = browser_export.build_preview_html(cards)
        del cards

        artifact_id = uuid.uuid4().hex
        artifact_tmpdir = tempfile.mkdtemp(prefix=f"store-card-{artifact_id}-")
        artifact_dir = Path(artifact_tmpdir)
        html_path = artifact_dir / "print.html"
        preview_path = artifact_dir / "preview.html"
        pdf_path = artifact_dir / "badges.pdf"
        png_zip_path = artifact_dir / "badges.zip"

        html_path.write_text(html_content, encoding="utf-8")
        preview_path.write_text(preview_html, encoding="utf-8")
        browser_export.export_html_to_pdf(html_content, str(pdf_path))

        ARTIFACTS[artifact_id] = {
            "created_at": time.time(),
            "tmpdir": artifact_tmpdir,
            "total_cards": total_cards,
            "total_pages": total_pages,
            "warnings": warnings,
            "html_path": str(html_path),
            "preview_path": str(preview_path),
            "pdf_path": str(pdf_path),
            "png_zip_path": str(png_zip_path),
            "png_zip_ready": False,
        }

        TMP_PREVIEW_PATH.write_text(preview_html, encoding="utf-8")
        del file_bytes, html_content, preview_html

        return jsonify(
            {
                "ok": True,
                "artifactId": artifact_id,
                "totalCards": total_cards,
                "totalPages": total_pages,
                "warnings": warnings,
                "employees": client_employees,
                "previewUrl": url_for("preview_artifact", artifact_id=artifact_id),
                "printUrl": url_for("print_artifact", artifact_id=artifact_id),
                "downloadHtmlUrl": url_for("download_html", artifact_id=artifact_id),
                "downloadPdfUrl": url_for("download_pdf", artifact_id=artifact_id),
                "downloadPngUrl": url_for("download_png_zip", artifact_id=artifact_id),
            }
        )
    except ValueError as exc:
        if artifact_tmpdir:
            shutil.rmtree(artifact_tmpdir, ignore_errors=True)
        gc.collect()
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception as exc:
        if artifact_tmpdir:
            shutil.rmtree(artifact_tmpdir, ignore_errors=True)
        gc.collect()
        return jsonify({"ok": False, "error": str(exc)}), 500
    finally:
        gc.collect()


@app.get("/preview/<artifact_id>")
def preview_artifact(artifact_id: str):
    try:
        artifact = get_artifact(artifact_id)
    except KeyError as exc:
        return Response(str(exc), status=404, mimetype="text/plain")
    return send_file(artifact["preview_path"], mimetype="text/html; charset=utf-8")


@app.get("/print/<artifact_id>")
def print_artifact(artifact_id: str):
    try:
        artifact = get_artifact(artifact_id)
    except KeyError as exc:
        return Response(str(exc), status=404, mimetype="text/plain")
    html = Path(artifact["html_path"]).read_text(encoding="utf-8").replace(
        "</body>",
        "<script>window.onload=function(){setTimeout(function(){window.print();},120);};</script></body>",
    )
    return Response(html, mimetype="text/html; charset=utf-8")


@app.get("/download/html/<artifact_id>")
def download_html(artifact_id: str):
    artifact = get_artifact(artifact_id)
    return send_file(
        artifact["preview_path"],
        as_attachment=True,
        download_name="store-card-latest-preview.html",
        mimetype="text/html",
    )


@app.get("/download/pdf/<artifact_id>")
def download_pdf(artifact_id: str):
    artifact = get_artifact(artifact_id)
    return send_file(
        artifact["pdf_path"],
        as_attachment=True,
        download_name="挂卡打印.pdf",
        mimetype="application/pdf",
    )


@app.get("/download/png/<artifact_id>")
def download_png_zip(artifact_id: str):
    artifact = get_artifact(artifact_id)
    if not artifact["png_zip_ready"]:
        build_png_zip(Path(artifact["pdf_path"]), artifact["total_pages"], Path(artifact["png_zip_path"]))
        artifact["png_zip_ready"] = True
    return send_file(
        artifact["png_zip_path"],
        as_attachment=True,
        download_name="挂卡打印.zip",
        mimetype="application/zip",
    )


@app.get("/latest-preview")
def latest_preview():
    if TMP_PREVIEW_PATH.exists():
        return redirect(TMP_PREVIEW_PATH.as_uri())
    return Response("尚未生成预览，请先上传 Excel 并生成工牌。", mimetype="text/plain; charset=utf-8")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8502, debug=False)
